"""Tests for Tier 2: golden-set schema, review-pool generator, and importer.

Kept separate from test_maestro.py (Tier 1's own test file, scoped to the check harness) since
this exercises a structurally different surface - file I/O, xlsx cell inspection, shuffle
determinism - with a different dependency (openpyxl, not lxml).
"""

from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

from maestro.golden.generate_pool import build_review_pool, shuffled_order
from maestro.golden.import_pool import import_pool, read_graded_rows
from maestro.golden.models import (
    GoldenExample,
    SmeVerification,
    append_jsonl,
    dump_golden_example,
    load_golden_example,
    read_jsonl,
    write_jsonl,
)
from maestro.golden.sme_roster import SmeAssignment, load_sme_roster

ROOT = Path(__file__).resolve().parents[1]


def _clean_candidate(example_id: str, exam_bank: str = "USMLE") -> GoldenExample:
    return GoldenExample(
        example_id=example_id, source_type="topic", exam_bank=exam_bank,
        question_type="single question",
        input={"topic": "test topic"},
        expected={
            "unique_name": f"{example_id}_unique",
            "question_text": "<p>Stem.</p>",
            "explanation_header": "<p>Answer.</p>",
            "explanation_footer": "<p>Footer.</p>",
            "main_topic": "Cardiology", "modifier": "Adult",
            "question_type": "single question", "question_format": "Text",
            # 3 references: satisfies references_format's reference_count sub-check (new since
            # the Tier 1 checks task filled that stub in) - keeps this fixture actually clean.
            "references": [
                "Smith J. Journal A. 2022.", "Doe R. Journal B. 2021.", "Lee K. Journal C. 2020.",
            ],
        },
        tags=["edge_case"],
    )


def _broken_candidate(example_id: str, exam_bank: str = "USMLE") -> GoldenExample:
    """A candidate whose 'expected' fails Tier 1's required_fields_present check."""
    candidate = _clean_candidate(example_id, exam_bank)
    candidate.expected["question_text"] = ""
    return candidate


class GoldenExampleSchemaTests(unittest.TestCase):
    def test_round_trip_preserves_sme_verification(self):
        """Specifically catches the naive GoldenExample(**data) regression: sme_verified must
        come back as a real SmeVerification instance, not a dict."""
        original = _clean_candidate("ex-1")
        original.sme_verified = SmeVerification(verified=True, reviewer_id="jdoe", verified_at="2026-01-01")

        round_tripped = load_golden_example(json.loads(json.dumps(dump_golden_example(original))))

        self.assertIsInstance(round_tripped.sme_verified, SmeVerification)
        self.assertEqual(round_tripped.sme_verified.reviewer_id, "jdoe")
        self.assertEqual(round_tripped, original)

    def test_defaults_are_unconfirmed_pending(self):
        example = GoldenExample()
        self.assertEqual(example.review_status, "pending_sme_review")
        self.assertFalse(example.sme_verified.verified)
        self.assertIsNone(example.sme_verified.reviewer_id)


class JsonlIoTests(unittest.TestCase):
    def test_read_missing_file_returns_empty_list(self):
        self.assertEqual(read_jsonl(Path("/nonexistent/path/does-not-exist.jsonl")), [])

    def test_write_read_round_trip(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "bank.jsonl"
            examples = [_clean_candidate("ex-1"), _clean_candidate("ex-2")]

            write_jsonl(path, examples)
            self.assertEqual(read_jsonl(path), examples)

    def test_append_adds_without_overwriting_existing(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "bank.jsonl"
            write_jsonl(path, [_clean_candidate("ex-1")])

            append_jsonl(path, [_clean_candidate("ex-2")])

            self.assertEqual([e.example_id for e in read_jsonl(path)], ["ex-1", "ex-2"])


class SmeRosterTests(unittest.TestCase):
    def test_empty_roster_returns_unconfirmed_assignment(self):
        roster = load_sme_roster({})
        assignment = roster.for_bank("USMLE")
        self.assertEqual(assignment, SmeAssignment())
        self.assertIsNone(assignment.reviewer_id)

    def test_unknown_bank_also_returns_unconfirmed(self):
        roster = load_sme_roster({"USMLE": {"reviewer_id": "jdoe"}})
        self.assertIsNone(roster.for_bank("COMLEX").reviewer_id)
        self.assertEqual(roster.for_bank("USMLE").reviewer_id, "jdoe")


class ReviewPoolShuffleTests(unittest.TestCase):
    def test_same_seed_is_deterministic(self):
        self.assertEqual(shuffled_order(10, seed=42), shuffled_order(10, seed=42))

    def test_different_seed_differs(self):
        self.assertNotEqual(shuffled_order(20, seed=1), shuffled_order(20, seed=2))


class ReviewPoolGeneratorTests(unittest.TestCase):
    def test_generated_workbook_has_expected_protection_and_validation(self):
        candidates = [_clean_candidate("ex-1"), _clean_candidate("ex-2")]

        workbook, mapping = build_review_pool(candidates, seed=1)

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "pool.xlsx"
            workbook.save(path)
            reloaded = openpyxl.load_workbook(path)
            sheet = reloaded.active

            self.assertEqual(sheet.freeze_panes, "A2")
            self.assertTrue(sheet.protection.sheet)

            headers = [cell.value for cell in sheet[1]]
            for name in ("grade", "reviewer_id", "notes"):
                column = headers.index(name) + 1
                self.assertFalse(sheet.cell(row=2, column=column).protection.locked, name)
            self.assertTrue(sheet.cell(row=2, column=headers.index("source_type") + 1).protection.locked)

            validations = sheet.data_validations.dataValidation
            self.assertTrue(any("approve" in (v.formula1 or "") for v in validations))

    def test_mapping_has_one_entry_per_candidate_and_no_identity_leak_in_sheet(self):
        candidates = [_clean_candidate("ex-1"), _clean_candidate("ex-2")]

        workbook, mapping = build_review_pool(candidates, seed=1)

        self.assertEqual(len(mapping["rows"]), 2)
        self.assertEqual({row["example_id"] for row in mapping["rows"].values()}, {"ex-1", "ex-2"})

        sheet = workbook.active
        headers = [cell.value for cell in sheet[1]]
        self.assertNotIn("example_id", headers)
        self.assertNotIn("ac_ref", headers)
        self.assertNotIn("unique_name", headers)


class ReviewPoolImporterTests(unittest.TestCase):
    def _grade_and_load(self, candidates, seed, grades):
        """grades: dict[example_id, (grade, reviewer_id)]. Returns (mapping, graded_rows)."""
        workbook, mapping = build_review_pool(candidates, seed=seed)
        sheet = workbook.active
        headers = [cell.value for cell in sheet[1]]
        display_by_example = {row["example_id"]: display_id for display_id, row in mapping["rows"].items()}
        grade_col, reviewer_col = headers.index("grade") + 1, headers.index("reviewer_id") + 1
        display_col = headers.index("display_id") + 1

        for row in range(2, sheet.max_row + 1):
            display_id = sheet.cell(row=row, column=display_col).value
            example_id = next(eid for eid, did in display_by_example.items() if did == display_id)
            if example_id in grades:
                grade, reviewer_id = grades[example_id]
                sheet.cell(row=row, column=grade_col).value = grade
                sheet.cell(row=row, column=reviewer_col).value = reviewer_id

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "graded.xlsx"
            workbook.save(path)
            graded_rows = read_graded_rows(path)
        return mapping, graded_rows

    def test_approved_row_is_promoted_with_correct_fields(self):
        candidate = _clean_candidate("ex-approve")
        mapping, graded_rows = self._grade_and_load([candidate], seed=1, grades={"ex-approve": ("approve", "jdoe")})

        with tempfile.TemporaryDirectory() as tmp:
            golden_dir = Path(tmp) / "golden"
            outcomes = import_pool(mapping, graded_rows, golden_dir=golden_dir, blocked_dir=Path(tmp))

            self.assertEqual(len(outcomes), 1)
            self.assertTrue(outcomes[0].promoted)

            promoted = read_jsonl(golden_dir / "usmle.jsonl")
            self.assertEqual(len(promoted), 1)
            self.assertEqual(promoted[0].review_status, "approved")
            self.assertTrue(promoted[0].sme_verified.verified)
            self.assertEqual(promoted[0].sme_verified.reviewer_id, "jdoe")
            self.assertIsNotNone(promoted[0].sme_verified.verified_at)

    def test_reject_and_needs_revision_are_blocked_not_promoted(self):
        candidates = [_clean_candidate("ex-reject"), _clean_candidate("ex-revise")]
        mapping, graded_rows = self._grade_and_load(
            candidates, seed=2,
            grades={"ex-reject": ("reject", "jdoe"), "ex-revise": ("needs_revision", "jdoe")},
        )

        with tempfile.TemporaryDirectory() as tmp:
            golden_dir, blocked_dir = Path(tmp) / "golden", Path(tmp) / "pools"
            outcomes = import_pool(mapping, graded_rows, golden_dir=golden_dir, blocked_dir=blocked_dir)

            self.assertTrue(all(not outcome.promoted for outcome in outcomes))
            self.assertFalse((golden_dir / "usmle.jsonl").exists())

            blocked_lines = (blocked_dir / "usmle.blocked.jsonl").read_text(encoding="utf-8").splitlines()
            reasons = {json.loads(line)["example_id"]: json.loads(line)["blocked_reason"] for line in blocked_lines}
            self.assertEqual(reasons["ex-reject"], "sme_rejected")
            self.assertEqual(reasons["ex-revise"], "sme_needs_revision")

    def test_tier1_failure_blocks_promotion_even_when_sme_approved(self):
        clean, broken = _clean_candidate("ex-clean"), _broken_candidate("ex-broken")
        mapping, graded_rows = self._grade_and_load(
            [clean, broken], seed=3,
            grades={"ex-clean": ("approve", "jdoe"), "ex-broken": ("approve", "jdoe")},
        )

        with tempfile.TemporaryDirectory() as tmp:
            golden_dir, blocked_dir = Path(tmp) / "golden", Path(tmp) / "pools"
            outcomes = import_pool(mapping, graded_rows, golden_dir=golden_dir, blocked_dir=blocked_dir)

            outcomes_by_id = {outcome.example_id: outcome for outcome in outcomes}
            self.assertTrue(outcomes_by_id["ex-clean"].promoted)
            self.assertFalse(outcomes_by_id["ex-broken"].promoted)
            self.assertEqual(outcomes_by_id["ex-broken"].reason, "tier1_check_failed")

            promoted = read_jsonl(golden_dir / "usmle.jsonl")
            self.assertEqual([e.example_id for e in promoted], ["ex-clean"])

    def test_no_run_checks_bypasses_the_tier1_gate(self):
        broken = _broken_candidate("ex-broken")
        mapping, graded_rows = self._grade_and_load([broken], seed=4, grades={"ex-broken": ("approve", "jdoe")})

        with tempfile.TemporaryDirectory() as tmp:
            golden_dir = Path(tmp) / "golden"
            outcomes = import_pool(
                mapping, graded_rows, golden_dir=golden_dir, blocked_dir=Path(tmp),
                run_tier1_gate=False,
            )
            self.assertTrue(outcomes[0].promoted)

    def test_collision_is_skipped_by_default_and_reported(self):
        candidate = _clean_candidate("ex-dup")
        mapping, graded_rows = self._grade_and_load([candidate], seed=5, grades={"ex-dup": ("approve", "jdoe")})

        with tempfile.TemporaryDirectory() as tmp:
            golden_dir = Path(tmp) / "golden"
            import_pool(mapping, graded_rows, golden_dir=golden_dir, blocked_dir=Path(tmp))
            second_pass = import_pool(mapping, graded_rows, golden_dir=golden_dir, blocked_dir=Path(tmp))

            self.assertFalse(second_pass[0].promoted)
            self.assertEqual(second_pass[0].reason, "collision")
            self.assertEqual(len(read_jsonl(golden_dir / "usmle.jsonl")), 1)

    def test_collision_replaces_in_place_with_allow_overwrite(self):
        candidate = _clean_candidate("ex-dup")
        mapping, graded_rows = self._grade_and_load([candidate], seed=6, grades={"ex-dup": ("approve", "jdoe")})

        with tempfile.TemporaryDirectory() as tmp:
            golden_dir = Path(tmp) / "golden"
            import_pool(mapping, graded_rows, golden_dir=golden_dir, blocked_dir=Path(tmp))
            second_pass = import_pool(
                mapping, graded_rows, golden_dir=golden_dir, blocked_dir=Path(tmp), allow_overwrite=True,
            )

            self.assertTrue(second_pass[0].promoted)
            self.assertEqual(len(read_jsonl(golden_dir / "usmle.jsonl")), 1)

    def test_ungraded_row_is_reported_and_not_promoted(self):
        candidate = _clean_candidate("ex-ungraded")
        mapping, graded_rows = self._grade_and_load([candidate], seed=7, grades={})

        with tempfile.TemporaryDirectory() as tmp:
            outcomes = import_pool(
                mapping, graded_rows, golden_dir=Path(tmp) / "golden", blocked_dir=Path(tmp),
            )
            self.assertFalse(outcomes[0].promoted)
            self.assertEqual(outcomes[0].reason, "ungraded")


class RunnerCliSmokeTests(unittest.TestCase):
    def test_generate_pool_help_is_available(self):
        result = subprocess.run(
            [sys.executable, str(ROOT / "maestro" / "golden" / "generate_pool.py"), "--help"],
            capture_output=True, text=True,
        )
        self.assertEqual(result.returncode, 0)
        self.assertIn("usage", result.stdout.lower())

    def test_import_pool_help_is_available(self):
        result = subprocess.run(
            [sys.executable, str(ROOT / "maestro" / "golden" / "import_pool.py"), "--help"],
            capture_output=True, text=True,
        )
        self.assertEqual(result.returncode, 0)
        self.assertIn("usage", result.stdout.lower())


if __name__ == "__main__":
    unittest.main()
